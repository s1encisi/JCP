import os
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── 全局 HV 参考点（内联定义，无需外部文件）───────────────────────────────────
# 目标顺序: [Cu_out, -As_out, E_total, -R_profit]（全部最小化方向）
# 参考点位于所有算法可能探索到的最差物理边界之外
GLOBAL_HV_REF_POINT = np.array([
    12.0,    # Cu_out：物理上限约 8，给足余量
    -0.0,    # -As_out：As 不可能为负，取 -0.0 确保涵盖所有负值
    1.0e6,   # E_total：极大耗电量上界
    1.0e6,   # -R_profit：极大亏损上界
])

# ── Excel 样式和工具函数 ───────────────────────────────────────────────────────

def _make_xl_styles() -> dict:
    thin = Side(style='thin')
    return {
        'hf':    Font(name='Arial', bold=True, size=10),
        'df':    Font(name='Arial', size=10),
        'ca':    Alignment(horizontal='center', vertical='center'),
        'la':    Alignment(horizontal='left',   vertical='center'),
        'bd':    Border(left=thin, right=thin, top=thin, bottom=thin),
        'hfill': PatternFill('solid', start_color='D9E1F2'),
    }

def _xl_header(ws, row: int, ncols: int, st: dict):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = st['hf']; cell.alignment = st['ca']
        cell.fill = st['hfill']; cell.border = st['bd']

def _xl_data(cell, val, st: dict, is_str: bool = False):
    cell.font = st['df']
    cell.alignment = st['la'] if is_str else st['ca']
    cell.border = st['bd']
    if val is None:
        return
    if isinstance(val, float):
        if not np.isnan(val):
            cell.value = round(val, 6)
    else:
        cell.value = val

# 模板各 sheet 的标准列头定义（统一在此处维护，消除三份副本的偏差）
_SHEET_HEADERS = {
    'Condition 1 Pareto Front':
        ['Cu_in', 'TA', 'IA', 'QA', 't',
         'Cu_out', 'As_out', 'E_total',
         'Net profit\n(Ten thousand  CNY)', 'VA', 'costs'],
    'Condition 2 Pareto Front':
        ['Cu_in', 'TB', 'IB', 'QB', 't',
         'Cu_out', 'As_out', 'E_total',
         'Net profit\n(Ten thousand  CNY)', 'VB', 'costs'],
    'Condition 3 Pareto Front':
        ['Cu_in', 'TA', 'IA', 'QA', 'TB', 'IB', 'QB', 't',
         'Cu_out', 'As_out', 'E_total',
         'Net profit\n(Ten thousand  CNY)', 'Vaverage', 'costs'],
    'Compromise Solutions':
        ['Condition', 'Scenario', 'Cu_in',
         'TA', 'IA', 'QA', 'TB', 'IB', 'QB', 't',
         'Cu_out', 'As_out', 'E_total', 'Net profit'],
    'rl_metrics':
        ['Condition', 'N_solutions', 'Hypervolume', 'Spacing'],
    'reward':
        ['Update_step', 'Condition 1', 'Condition 2', 'Condition 3'],
    'quality_metrics_ppo':
        ['Condition', 'N_solutions',
         'Mean_Cu_out', 'Std_Cu_out', 'Range_Cu_out',
         'Mean_As_out', 'Std_As_out', 'Range_As_out',
         'Mean_E_total', 'Std_E_total', 'Range_E_total',
         'Mean_Net_profit', 'Std_Net_profit', 'Range_Net_profit'],
}

# ── 合并 RL 结果的主函数 ───────────────────────────────────────────────────────

def merge_rl_results(base_dir: str, final_output: str):
    """
    将 Condition 1/2/3 的 pareto_ppo_condition 文件合并为完整的 RL_final.xlsx。
    前提：pareto_ppo_condition{1,2,3}.csv 已存在。
    """

    # 决策变量映射
    dec_var_map = {
        'Condition 1': ['Cu_in', 'TA', 'IA', 'QA', 't'],
        'Condition 2': ['Cu_in', 'TB', 'IB', 'QB', 't'],
        'Condition 3': ['Cu_in', 'TA', 'IA', 'QA', 'TB', 'IB', 'QB', 't'],
    }

    voltage_col_map = {
        'Condition 1': 'VA',
        'Condition 2': 'VB',
        'Condition 3': 'Vaverage',
    }

    files = {
        'Condition 1': os.path.join(base_dir, 'outputs_condition1fast', 'pareto_ppo_condition1.csv'),
        'Condition 2': os.path.join(base_dir, 'outputs_condition2fast', 'pareto_ppo_condition2.csv'),
        'Condition 3': os.path.join(base_dir, 'outputs_condition3fast', 'pareto_ppo_condition3.csv'),
    }
    for cond, path in files.items():
        if not os.path.exists(path):
            raise FileNotFoundError(f"[merge] 找不到 {cond} 文件: {path}")

    print("=== 开始合并 RL 结果 ===")
    sheet_names = list(_SHEET_HEADERS.keys())
    
    # 读取所有文件数据
    dfs = {}
    for cond, path in files.items():
        print(f"  读取 {cond} ...")
        df = pd.read_csv(path)
        dfs[cond] = df

    merged = {}
    
    # ── A. 处理各工况的 Pareto Front ─────────────────────────────────────────
    for cond in files:
        df = dfs[cond].copy()
        # 选择需要的列
        dec_vars = dec_var_map[cond]
        obj_cols = ['Cu_out', 'As_out', 'E_total', 'Net_profit', 'V', 'costs_max']
        selected_cols = dec_vars + obj_cols
        pareto_df = df[selected_cols]
        # 重命名列以匹配模板
        rename_map = {'V': voltage_col_map[cond], 'costs_max': 'costs', 'Net_profit': 'Net profit\n(Ten thousand  CNY)'}
        pareto_df = pareto_df.rename(columns=rename_map)
        merged[f'{cond} Pareto Front'] = pareto_df

    # ── B. 生成 Compromise Solutions ───────────────────────────────────────
    comp_dfs = []
    for cond, df in dfs.items():
        if len(df) == 0:
            continue
        
        # 计算归一化目标矩阵
        obj_mat = np.column_stack([
            df['Cu_out'].values,
            -df['As_out'].values,
            df['E_total'].values,
            -df['Net_profit'].values,
        ])
        mn = obj_mat.min(axis=0); mx = obj_mat.max(axis=0)
        dn = mx - mn; dn[dn == 0] = 1e-10
        obj_norm = (obj_mat - mn) / dn
        
        # 选择不同场景的解
        scene_map = [
            ('Quality priority',  int(np.argmin(obj_norm[:, 0] + obj_norm[:, 1]))),
            ('Energy priority',   int(np.argmin(obj_norm[:, 2]))),
            ('Economic priority', int(np.argmin(obj_norm[:, 3]))),
            ('Balanced',          int(np.argmin(obj_norm.sum(axis=1)))),
        ]
        
        for scenario, row_idx in scene_map:
            rr = df.iloc[row_idx]
            row_data = {
                'Condition': cond,
                'Scenario': scenario,
                'Cu_in': float(rr['Cu_in']),
                'TA': float(rr.get('TA', 0)) if 'TA' in rr else None,
                'IA': float(rr.get('IA', 0)) if 'IA' in rr else None,
                'QA': float(rr.get('QA', 0)) if 'QA' in rr else None,
                'TB': float(rr.get('TB', 0)) if 'TB' in rr else None,
                'IB': float(rr.get('IB', 0)) if 'IB' in rr else None,
                'QB': float(rr.get('QB', 0)) if 'QB' in rr else None,
                't': float(rr['t']),
                'Cu_out': float(rr['Cu_out']),
                'As_out': float(rr['As_out']),
                'E_total': float(rr['E_total']),
                'Net profit': float(rr['Net_profit']),
            }
            comp_dfs.append(row_data)

    merged['Compromise Solutions'] = pd.DataFrame(comp_dfs) if comp_dfs else pd.DataFrame()

    # ── C. 生成 rl_metrics ─────────────────────────────────────────────────
    rl_metrics_data = []
    for cond, df in dfs.items():
        n_sol = len(df)
        # 计算超体积和间距
        hv_val = None
        sp_val = None
        
        try:
            from pymoo.indicators.hv import HV
            from pymoo.indicators.spacing import SpacingIndicator
            F_rl = np.column_stack([
                df['Cu_out'].values, -df['As_out'].values,
                df['E_total'].values, -df['Net_profit'].values,
            ])
            # 使用全局参考点
            hv_val = float(HV(ref_point=GLOBAL_HV_REF_POINT)(F_rl))
            sp_val = float(SpacingIndicator()(F_rl))
        except (ImportError, Exception) as e:
            print(f"  Warning: HV calculation failed for {cond}: {e}")
            pass
        
        rl_metrics_data.append({
            'Condition': cond,
            'N_solutions': n_sol,
            'Hypervolume': hv_val,
            'Spacing': sp_val,
        })

    merged['rl_metrics'] = pd.DataFrame(rl_metrics_data)

    # ── D. reward：从原来的Excel文件中读取 ───────────────────────────
    reward_files = {
        'Condition 1': os.path.join(base_dir, 'outputs_condition1fast', 'RL_results_condition1.xlsx'),
        'Condition 2': os.path.join(base_dir, 'outputs_condition2fast', 'RL_results_condition2.xlsx'),
        'Condition 3': os.path.join(base_dir, 'outputs_condition3fast', 'RL_results_condition3.xlsx'),
    }
    reward_dfs = []
    for cond, path in reward_files.items():
        if os.path.exists(path):
            try:
                xl = pd.ExcelFile(path)
                if 'reward' in xl.sheet_names:
                    df = pd.read_excel(path, sheet_name='reward')
                    if not df.empty:
                        # 重命名列以匹配模板
                        if 'reward' in df.columns:
                            df = df.rename(columns={'reward': cond})
                        reward_dfs.append(df)
            except Exception as e:
                print(f"  Warning: Failed to read reward data from {cond}: {e}")
                pass
    
    if reward_dfs:
        # 以 Update_step 为键做外连接
        df_merged = reward_dfs[0]
        for df_next in reward_dfs[1:]:
            df_merged = pd.merge(df_merged, df_next, on='Update_step', how='outer')
        merged['reward'] = df_merged.sort_values('Update_step').reset_index(drop=True)
    else:
        merged['reward'] = pd.DataFrame()

    # ── E. 生成 quality_metrics_ppo ───────────────────────────────────────
    quality_data = []
    for cond, df in dfs.items():
        if len(df) == 0:
            continue
        
        quality_data.append({
            'Condition': cond,
            'N_solutions': len(df),
            'Mean_Cu_out': float(df['Cu_out'].mean()),
            'Std_Cu_out': float(df['Cu_out'].std()),
            'Range_Cu_out': float(df['Cu_out'].max() - df['Cu_out'].min()),
            'Mean_As_out': float(df['As_out'].mean()),
            'Std_As_out': float(df['As_out'].std()),
            'Range_As_out': float(df['As_out'].max() - df['As_out'].min()),
            'Mean_E_total': float(df['E_total'].mean()),
            'Std_E_total': float(df['E_total'].std()),
            'Range_E_total': float(df['E_total'].max() - df['E_total'].min()),
            'Mean_Net_profit': float(df['Net_profit'].mean()),
            'Std_Net_profit': float(df['Net_profit'].std()),
            'Range_Net_profit': float(df['Net_profit'].max() - df['Net_profit'].min()),
        })

    merged['quality_metrics_ppo'] = pd.DataFrame(quality_data) if quality_data else pd.DataFrame()

    print("\n  写入合并文件 ...")
    st  = _make_xl_styles()
    wb  = Workbook()
    del wb['Sheet']

    for sn in sheet_names:
        ws = wb.create_sheet(title=sn)
        df = merged[sn]
        if df.empty:
            print(f"  ⚠ 跳过 {sn}（无数据）")
            continue
        _xl_header(ws, 1, len(df.columns), st)
        # 写入列头
        headers = _SHEET_HEADERS[sn]
        for c_idx, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c_idx).value = h
        # 写入数据
        for r_idx, row in enumerate(df.itertuples(index=False), start=2):
            for c_idx, val in enumerate(row, start=1):
                _xl_data(ws.cell(row=r_idx, column=c_idx), val, st, is_str=(c_idx <= 2 and sn == 'Compromise Solutions'))
        # 调整列宽
        for c in range(1, len(df.columns) + 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 12

    wb.save(final_output)
    print(f"\n=== 合并完成 ===")
    print(f"  输出文件: {final_output}")

# ── 主函数调用示例 ───────────────────────────────────────────────────────────
if __name__ == "__main__":
    # Use the reorganized project copy only; never read from or write to the
    # original source tree implicitly.
    script_dir = Path(__file__).resolve().parent
    base_dir = script_dir / "result"
    final_output = script_dir / "RL_final.xlsx"
    
    try:
        merge_rl_results(str(base_dir), str(final_output))
    except Exception as e:
        print(f"合并失败: {e}")
